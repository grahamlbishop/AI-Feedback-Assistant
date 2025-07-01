import os
import google.generativeai as genai
from docx import Document  # For reading .docx
import PyPDF2  # For reading .pdf
import time  # For rate limiting
import sys  # To exit cleanly on error

# --- Configuration ---

# === IMPORTANT: Secure your API Key! ===
# Reads the API key from an environment variable named GOOGLE_API_KEY
api_key = os.environ.get("GOOGLE_API_KEY")
if not api_key:
    print("Error: GOOGLE_API_KEY environment variable not found.")
    print("Please set the environment variable before running the script.")
    print("Refer to instructions on how to set environment variables for your OS.")
    print("Exiting.")
    sys.exit(1) # Exit the script if key is not found

# Configure the Gemini API client
try:
    genai.configure(api_key=api_key)
except Exception as e:
    print(f"Error configuring the Google AI client: {e}")
    print("Please ensure your API key is valid.")
    sys.exit(1)

# === Model Configuration ===
# Consider later and/or pro models for potentially higher quality (check pricing/availability)
MODEL_NAME = 'gemini-2.0-flash' # Flash is faster and cheaper, Pro might be better quality
try:
    model = genai.GenerativeModel(MODEL_NAME)
    print(f"Using Generative Model: {MODEL_NAME}")
except Exception as e:
    print(f"Error initializing Generative Model ('{MODEL_NAME}'): {e}")
    print("Check if the model name is correct and your API key has access.")
    sys.exit(1)


# === Folder Paths === (Relative to where the script is run)
papers_folder = 'papers'
output_folder = 'feedback'
# Ensure input folder exists
if not os.path.isdir(papers_folder):
    print(f"Error: Input folder '{papers_folder}' not found.")
    print("Please create it and place the paper files inside.")
    sys.exit(1)
# Ensure output folder exists
os.makedirs(output_folder, exist_ok=True)


# === Assignment Context ===
assignment_description = """
Project 1: Analyzing a Text

This project asks you to analyze one of the following texts from class -- Dungy, Bastian, or Young. Different from a summary where you share with your reader the main arguments of something you read, an analysis asks you to interpret why the author crafted their piece in the manner that they did and how that piece addresses a broader conversation.

In class, we have been analyzing this texts by considering what each authors driving questions has been. What is the underlying question that pushes their thought and analysis forward? Your essay should have its own driving question (an alternative to a thesis statement) that motivates the analysis and argument of your essay writing and structure.

This project also provides an opportunity to work on the following learning goals:
- Engage with current conversations taking place around the climate crisis to develop critical reading and thinking skills
- Reexamine and refine the processes of one’s own writing to enhance self-awareness as a writer and develop strategies for effectively communicating with diverse audiences.
- Learn about and begin to develop ethical communication processes
- Learn about the role of collaboration in the learning process
- Reflect on the learning environment of the course and engage in co-creating learning experiences

In order to complete this project you need to follow these steps:
Step 1: Choose the text you wish to analyze (Dungy, Bastian, or Young)
Step 2: Decide how you will analyze the text. Think about some of the questions discussed in class:
    - What is the driving question that the author is trying to answer? What sub questions do they attempt to answer? Does the author identify a problem and/or propose a solution? What is the purpose of discussing these questions/solutions?
    - Who are they posing this question to? Who is their audience? How can you tell? How do they appeal to their audience's values/knowledge/conventions?
    - Who is the author? How does their intersectionality inform the text?
    - What is the author’s positionality in relation to their driving question and/or audience? How does this positionality impact how they discuss their driving question?
    - How does the author communicate their argument?  (You might consider using rhetorical devices such as humor,  or observe other strategies they use to craft their message.)
    - What other authors/sources do they draw on as evidence to support their argument? How effectively do they build their case?
You can focus on one of these questions to create your own driving question about the text.
Step 3: Write a 3-4 page first draft of your analysis. Your analysis should include an argument or driving question in response to the question you chose to use to analyze the text. First focus on your analysis, and then think about the best way to present/craft the analysis in your text. You do not need to use a standard 5-paragraph essay.
Step 4: Submit your first draft to your instructor, and be sure to also complete your pre-conference reflection.
Step 5: Meet for a mid-term conference with your instructor to discuss your analysis.
Step 6: After meeting with your instructor, revise your project. Submit your revision within a week after your conference. When submitting your revision, please include: The first draft, instructor comments, revised draft, and a pre-conference reflection.
"""

# === Example Feedback ===
# Provides style, tone, and specificity guidance. AI should adapt content.
example_feedback_letter = """
Dear student,
Thank you for sharing this essay with me. You’ve effectively identified Dungy’s core argument, and your use of evidence (Dungy’s childhood anecdotes, her discussion of ecopoetics, her rebuttal to counterarguments) is relevant to your interest in “different definitions of an environment.”
For revision, I would focus on refining your core argument about Dungy’s paper. Your driving question currently reads like a rephrasing of Dungy’s, so I might suggest returning to your initial interest in “different definitions of an environment” and see if you can unpack the how behind Dungy’s argument. What literary strategies does she employ to capture these “different definitions”? How might you structure your paper in a way that allows you to compare and contrast these definitions that Dungy explores more clearly? Pay particular attention to the transitions between sections, and focus on telling a story for your reader.
I would also suggest reading your essay aloud and proofreading it for clarity and flow. Certain word choices and typos (for example, “pressing” in the paper’s opening sentence, pronoun consistency, or “Ddungy” instead of “Dungy”) got in the way of my reading experience, so make sure every sentence conveys what you want it to say.
Please let me know if you have any questions about my comments. I look forward to exploring the essay in more depth with you during the conference.
"""

# === Feedback Template (Revised) ===
# The AI will fill in the bracketed parts based on its analysis.
# Remember you will manually replace the {student_identifier} with the actual name later.
feedback_template = """
Dear {student_identifier},

Thank you for submitting your Project 1 analysis. Here is your feedback:

**Overall Analysis & Argument:**
*   [AI: Comment on the clarity and insightfulness of the student's analysis of the chosen text (Dungy, Bastian, or Young). Does the essay move beyond summary? Does it seem guided by its own clear driving question or interpretive thesis, as requested in the assignment?]
*   [AI: Identify a key strength in the analysis or argument, citing a specific example or quote from the student's essay.]

**Use of Evidence & Examples:**
*   [AI: Comment on how effectively the student uses evidence (quotes, specific examples) from the primary text (Dungy, Bastian, or Young) to support their analytical points. Are the examples well-chosen, integrated smoothly, and explained in relation to the student's argument?]
*   [AI: Note a strength or area for improvement regarding evidence use, citing a specific example or quote from the student's essay.]

**Structure & Clarity:**
*   [AI: Comment on the overall organization and flow of the essay. Is the argument logical and easy to follow? Are transitions between paragraphs/ideas effective? Does the structure support the analysis? Consider the learning goals about communicating effectively.]
*   [AI: Note a strength or area for improvement regarding structure or clarity, citing a specific example or quote from the student's essay.]

**Areas for Revision Focus:**
*   [AI: Based on the assignment goals (analysis vs summary, driving question, learning goals) and the analysis of this specific paper, suggest 1-2 concrete, actionable areas for the student to focus on during revision. Try to connect suggestions back to the assignment description.]
*   [AI: Suggest specific strategies, questions for the student to consider, or parts of their essay to revisit during revision, aiming for the constructive and specific style shown in the example feedback.]

**Proofreading Note (Optional but Recommended):**
*   [AI: Briefly mention if consistent patterns of errors (grammar, typos, awkward phrasing, citation issues) significantly impeded readability. If so, gently suggest careful proofreading, perhaps reading aloud as mentioned in the example feedback.]

I look forward to discussing your analysis further during our conference. Please come prepared with any questions you have about this feedback or your revision ideas.

Best regards,
[Your Name]
"""

# === Base Prompt for the AI ===
# This combines all instructions, context, and placeholders.
base_prompt = f"""
You are an AI teaching assistant providing feedback on a student's Project 1 analysis paper draft.
Your goal is to generate helpful, specific, and constructive feedback to guide the student's revision process, aligning with the assignment's goals and the instructor's desired feedback style.

**Assignment Context:**
This is the description for Project 1 the student was given. Pay close attention to the key requirements: analyzing (not summarizing) Dungy, Bastian, or Young; developing their *own* driving question; and the listed learning goals.
{assignment_description}

**Example of Desired Feedback:**
This example shows the desired tone, style, and level of specificity (connecting comments to the text, offering concrete revision suggestions). Adapt the *content* to the student paper you analyze, but match the *style*. Note the example discusses Dungy; the current paper might analyze Bastian or Young.
{example_feedback_letter}

**Instructions for Generating Feedback:**
1.  Thoroughly analyze the **Student Paper Text** provided below.
2.  Identify which primary text (Dungy, Bastian, or Young) the student is analyzing.
3.  Evaluate the paper based on the **Assignment Context**, focusing on analysis depth, the presence and clarity of the student's driving question/thesis, use of evidence, and structure.
4.  Generate a feedback letter using the **Feedback Template** provided below.
5.  **Adhere strictly to the structure** of the Feedback Template, filling in each bracketed `[AI: ...]` section.
6.  Emulate the **constructive tone and specificity** shown in the **Example Feedback**. **Crucially, use specific examples or short quotes from the student's paper text** to justify your points in each section.
7.  The student's identifier (likely their username) is provided. Use this identifier *only* in the salutation `Dear {{student_identifier}},` as shown in the template.

**Feedback Template:**
{feedback_template}

**Student Identifier:** {{student_identifier_placeholder}}

**Student Paper Text:**
{{paper_text_placeholder}}

**Generate the feedback letter now, following all instructions carefully:**
"""

# --- Processing Loop ---
print(f"\n--- Starting Batch Feedback Generation ---")

try:
    # List files, ignore hidden files/folders
    paper_files = [f for f in os.listdir(papers_folder) if os.path.isfile(os.path.join(papers_folder, f)) and not f.startswith('.')]
except FileNotFoundError:
    print(f"Error: Input folder '{papers_folder}' not found.")
    sys.exit(1)

if not paper_files:
    print(f"No files found in the '{papers_folder}' folder.")
    sys.exit(0)

total_files = len(paper_files)
processed_count = 0
error_count = 0

print(f"Found {total_files} files to process in '{papers_folder}'. Outputting to '{output_folder}'.")

for index, filename in enumerate(paper_files):
    print("-" * 50) # Separator for clarity
    print(f"Processing file {index + 1}/{total_files}: {filename}")

    full_text = ""
    student_identifier = "Unknown_Student" # Default identifier

    # --- Extract Student Identifier (Handles username_ID_ID_OriginalName.ext) ---
    try:
        base_name = os.path.splitext(filename)[0] # Remove extension
        parts = base_name.split('_', 1) # Split only at the first underscore
        if parts and parts[0].strip():
            student_identifier = parts[0].strip()
        elif base_name.strip(): # Fallback if no underscore
             student_identifier = base_name.strip()
        else:
             print(f"  Warning: Could not extract a valid identifier from filename. Using default.")
        print(f"  Extracted Identifier: {student_identifier}")
    except Exception as e:
        print(f"  Warning: Error extracting identifier from filename '{filename}'. Using default. Error: {e}")

    filepath = os.path.join(papers_folder, filename)
    file_processed = False # Flag for rate limiting

    # --- Read File Content based on extension ---
    try:
        if filename.lower().endswith(".docx"):
            print("  Reading DOCX file...")
            doc = Document(filepath)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            full_text = "\n\n".join(paragraphs) # Use double newline for better paragraph separation
            if not full_text:
                print(f"  Warning: No text extracted from DOCX file or file is empty.")
                error_count += 1
                continue

        elif filename.lower().endswith(".pdf"):
            print("  Reading PDF file...")
            text_list = []
            try:
                with open(filepath, 'rb') as pdf_file:
                    reader = PyPDF2.PdfReader(pdf_file)
                    # Check for encryption
                    if reader.is_encrypted:
                        try:
                            # Try decrypting with empty password (common case)
                            if reader.decrypt('') == PyPDF2.PasswordType.NOT_DECRYPTED:
                                print(f"  Warning: Skipping password-protected PDF: {filename}")
                                error_count += 1
                                continue # Skip this file
                        except Exception as decrypt_err:
                             print(f"  Warning: Skipping encrypted PDF (decryption failed): {filename} - {decrypt_err}")
                             error_count += 1
                             continue

                    num_pages = len(reader.pages)
                    # print(f"  Found {num_pages} page(s). Extracting text...") # Optional verbosity
                    for page_num in range(num_pages):
                        try:
                            page = reader.pages[page_num]
                            extracted = page.extract_text()
                            if extracted:
                                text_list.append(extracted.strip())
                        except Exception as page_error:
                             print(f"    Warning: Error extracting text from PDF page {page_num + 1}: {page_error}")
                full_text = "\n\n".join(text_list).strip() # Join pages with double newline
                if not full_text:
                    print(f"  Warning: No text extracted from PDF (check if image-based or complex).")
                    error_count += 1
                    continue
            except ImportError as ie:
                 print(f"  Error: PyPDF2 dependency possibly missing or corrupt: {ie}")
                 print(f"  Try: pip install --upgrade PyPDF2")
                 error_count +=1
                 continue
            except Exception as pdf_err:
                print(f"  Error reading PDF file structure {filename}: {pdf_err}")
                error_count += 1
                continue

        else:
            print(f"  Skipping unsupported file type: {filename}")
            # Not counted as an error, just skipped.
            continue

        # --- If text was extracted successfully ---
        if full_text:
            print(f"  Extracted text length: ~{len(full_text)} characters.")
            # --- Prepare the final prompt for the API ---
            try:
                # Replace ALL placeholders in the base prompt string
                prompt_for_api = base_prompt.replace("{assignment_context_placeholder}", assignment_description)
                prompt_for_api = prompt_for_api.replace("{example_feedback_placeholder}", example_feedback_letter)
                prompt_for_api = prompt_for_api.replace("{student_identifier_placeholder}", student_identifier)
                prompt_for_api = prompt_for_api.replace("{paper_text_placeholder}", full_text)
                # Replace placeholder in the template part itself
                prompt_for_api = prompt_for_api.replace("{student_identifier}", student_identifier)

                # --- Call the Google Gemini API ---
                print("  Sending request to Gemini API...")
                response = model.generate_content(
                    prompt_for_api,
                    # Optional: Add safety settings if needed, balancing safety and utility
                    # safety_settings=[
                    #     {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_LOW_AND_ABOVE"}, # Stricter
                    #     {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_LOW_AND_ABOVE"}, # Stricter
                    #     {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_MEDIUM_AND_ABOVE"},
                    #     {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_MEDIUM_AND_ABOVE"},
                    # ]
                    # Optional: Control generation parameters
                    # generation_config=genai.types.GenerationConfig(
                    #     # candidate_count=1, # Default is 1
                    #     # stop_sequences=['\n\n\n'], # Example stop sequence
                    #     # max_output_tokens=2048, # Limit output length
                    #     temperature=0.7 # 0.0 = deterministic, 1.0 = max creativity
                    # )
                )

                # --- Extract and Save Feedback ---
                feedback_text = ""
                try:
                    # Accessing response text safely - check candidate parts
                    if response.parts:
                         feedback_text = "".join(part.text for part in response.parts).strip()
                    else:
                         # Sometimes .text might work even if .parts is empty, try as fallback
                         feedback_text = response.text.strip()

                except AttributeError:
                     # Handle cases where .text might not exist if response failed early
                     print("  Warning: Could not directly access .text attribute in API response.")
                     feedback_text = "" # Ensure it's empty
                except Exception as resp_err:
                     print(f"  Warning: Could not extract text from API response parts. Error: {resp_err}")
                     feedback_text = "" # Ensure it's empty

                # Check if feedback is empty or blocked
                if feedback_text:
                    output_filename = os.path.join(output_folder, f"{student_identifier}_feedback.txt")
                    with open(output_filename, 'w', encoding='utf-8') as f:
                        f.write(feedback_text)
                    print(f"  Successfully generated and saved feedback to '{output_filename}'")
                    processed_count += 1
                    file_processed = True
                else:
                    # Handle blocked prompts or genuinely empty responses
                    print(f"  Warning: No feedback content generated for {filename}.")
                    try:
                        # Log safety feedback if available
                        block_reason = response.prompt_feedback.block_reason
                        safety_ratings = response.prompt_feedback.safety_ratings
                        print(f"    Block Reason (if any): {block_reason}")
                        print(f"    Safety Ratings: {safety_ratings}")
                        error_filename = os.path.join(output_folder, f"{student_identifier}_ERROR_FeedbackBlockedOrEmpty.txt")
                        with open(error_filename, 'w', encoding='utf-8') as f:
                            f.write(f"Feedback generation blocked or empty for {filename} (Identifier: {student_identifier}).\n")
                            f.write(f"Block Reason: {block_reason}\n")
                            f.write(f"Safety Ratings: {safety_ratings}\n")
                    except Exception:
                         print("    Could not retrieve detailed safety/block feedback from response.")
                         error_filename = os.path.join(output_folder, f"{student_identifier}_ERROR_EmptyResponse.txt")
                         with open(error_filename, 'w', encoding='utf-8') as f:
                            f.write(f"API returned an empty response for {filename} (Identifier: {student_identifier}).\n")
                    error_count += 1
                    continue # Skip saving/pausing

            except Exception as api_error:
                print(f"!! Error during API call or response processing for {filename}: {api_error}")
                # You might want to log the specific error to a file here too
                error_filename = os.path.join(output_folder, f"{student_identifier}_ERROR_API_Call_Failed.txt")
                with open(error_filename, 'w', encoding='utf-8') as f:
                   f.write(f"API call failed for {filename} (Identifier: {student_identifier}).\n")
                   f.write(f"Error: {api_error}\n")
                error_count += 1
                continue # Skip to next file

            # --- Rate Limiting ---
            if file_processed: # Only pause if we successfully processed and saved
                 print("  Pausing briefly to respect API rate limits...")
                 # Adjust sleep time based on your API tier limits and observation
                 # Free tiers often have low requests-per-minute limits (e.g., 15-60 RPM)
                 time.sleep(4) # Pause for 4 seconds (adjust if needed)

    except Exception as file_proc_error:
        print(f"!! Unexpected error processing file {filename} before API call: {file_proc_error}")
        error_filename = os.path.join(output_folder, f"{student_identifier}_ERROR_File_Processing.txt")
        with open(error_filename, 'w', encoding='utf-8') as f:
            f.write(f"Unexpected error processing file {filename} (Identifier: {student_identifier}) before API call.\n")
            f.write(f"Error: {file_proc_error}\n")
        error_count += 1
        continue # Skip to next file

# --- Final Summary ---
print("-" * 50)
print("\n--- Batch Feedback Generation Summary ---")
print(f"Total files found in '{papers_folder}': {total_files}")
print(f"Successfully generated feedback for: {processed_count} files")
print(f"Files skipped or resulting in errors: {error_count}")
print(f"Feedback files (and any error logs) saved in the '{output_folder}' folder.")
print("\n--- IMPORTANT REMINDERS ---")
print("1. REVIEW AND EDIT EACH feedback file carefully before sharing.")
print("2. Manually replace the '{student_identifier}' (username) in each feedback letter with the student's actual name.")
print("-" * 50)